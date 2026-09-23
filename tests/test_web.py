"""Offline HTTP boundary tests: real local catalogue, temporary state, no APIs."""

import asyncio
from copy import deepcopy
from dataclasses import replace
import html
from types import SimpleNamespace
from unittest.mock import AsyncMock

from fastapi.testclient import TestClient
import httpx
import pytest
from telegram.ext import CallbackQueryHandler, CommandHandler, MessageHandler

from app import main
from app.agent import Assistant
from app.catalog import Catalog
from app.config import ROOT, Settings
from app.telegram_bot import build_application


@pytest.fixture(scope="module")
def catalog():
    # Build the real search index only once; the fixture never refreshes EKT.
    return Catalog(ROOT / "data")


@pytest.fixture
def web(tmp_path, monkeypatch, catalog):
    config = Settings(api_key="", telegram_token="", bot_username="",
                      data_dir=tmp_path, db_path=tmp_path / "web.sqlite",
                      public_base_url="http://testserver")
    monkeypatch.setattr(main, "settings", config)
    monkeypatch.setattr(main, "Catalog", lambda _path: catalog)
    monkeypatch.setattr(main, "Assistant", lambda cat, store: Assistant(cat, store, config))
    async_network = AsyncMock(side_effect=AssertionError("External HTTP is forbidden in web tests"))
    monkeypatch.setattr(httpx.AsyncClient, "send", async_network)

    def reject_sync_network(*_args, **_kwargs):
        raise AssertionError("External HTTP is forbidden in web tests")

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", reject_sync_network)

    async def local_stock(product_id):
        product = catalog.get(product_id)
        product["stock_fresh"] = True
        return product

    refresh = AsyncMock(side_effect=local_stock)
    monkeypatch.setattr(catalog, "refresh", refresh)
    with TestClient(main.app, base_url=config.public_base_url) as client:
        state = SimpleNamespace(client=client, assistant=main.app.state.assistant,
                                config=config, refresh=refresh)
        yield state
    assert async_network.call_count == 0, "A test attempted an external network request"


def session_key(client):
    response = client.get("/api/cart")
    assert response.status_code == 200
    return "web:" + client.cookies.get("ekt_session")


def proposal_for(web, client, quantity=1, overrides=None):
    key = session_key(client)
    product = next(deepcopy(p) for p in web.assistant.catalog._items.values()
                   if p["quantity"] >= 10 and p["price"] > 0)
    item = {field: product[field] for field in ("id", "article", "name", "price")}
    item["quantity"] = quantity
    item.update(overrides or {})
    return web.assistant.store.propose(key, [item]), product


def test_health_and_index_are_available_without_ai_or_telegram(web, catalog):
    response = web.client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert data["catalog_products"] == catalog.count > 0
    assert data["openai_configured"] is False
    assert data["telegram"] == "not_configured"
    assert data["ekt_cart_bridge"] is False
    page = web.client.get("/")
    assert page.status_code == 200
    assert "text/html" in page.headers["content-type"]


def test_chat_existing_article_returns_real_catalogue_data(web, catalog):
    product = catalog.get("010500006_")
    assert product is not None, "Known EKT acceptance-test article is missing"
    response = web.client.post("/api/chat", json={"text": product["article"]})
    assert response.status_code == 200
    data = response.json()
    assert data["products"][0]["id"] == product["id"]
    assert data["products"][0]["price"] == product["price"]
    assert data["products"][0]["quantity"] == product["quantity"]
    assert product["article"] in data["text"]
    assert product["name"] in data["text"]
    assert data["trace"] == ["product_details"]
    assert data["usage"] == []
    assert data["products"][0]["stock_fresh"] is False
    assert "снимок каталога" in data["text"]
    assert "проверено в API" not in data["text"]
    assert data["proposal"] is None
    assert web.client.get("/api/cart").json()["items"] == []
    web.refresh.assert_not_awaited()


def test_exact_article_remains_local_even_after_previous_live_refresh(web, catalog, monkeypatch):
    product = catalog.get("010500006_")
    # A recent confirmation may have populated live fields in the shared cache.
    # Ordinary consultation must still present them as a locally stored snapshot.
    monkeypatch.setitem(catalog._items, product["id"], product | {"stock_fresh": True, "stock_source": "ekt_api"})
    web.refresh.side_effect = AssertionError("Ordinary product questions must never refresh EKT")
    response = web.client.post("/api/chat", json={"text": product["article"]})
    assert response.status_code == 200
    data = response.json()
    assert data["products"][0]["stock_fresh"] is False
    assert "снимок каталога" in data["text"]
    assert "проверено в API" not in data["text"]
    web.refresh.assert_not_awaited()


def test_product_details_tool_reads_only_local_catalogue(web, catalog):
    product = catalog.get("010500006_")
    web.refresh.side_effect = AssertionError("Product details must never refresh EKT")
    details = asyncio.run(web.assistant.run_tool("web:test", "product_details", {"product_id": product["id"]}))
    assert details["id"] == product["id"]
    assert details["price"] == product["price"]
    assert details["quantity"] == product["quantity"]
    assert details["stock_fresh"] is False
    web.refresh.assert_not_awaited()


def test_purchase_terms_answer_contains_all_three_topics_and_source(web):
    response = web.client.post("/api/chat", json={"text": "Какие оплата, доставка и минимальная партия?"})
    assert response.status_code == 200
    text = response.json()["text"]
    assert "Физлица" in text and "Юрлица" in text
    assert "30 000" in text and "менеджер" in text
    assert "упаковки" in text
    assert "https://ekt.kz/checkout-delivery/" in text
    assert web.client.get("/api/cart").json()["items"] == []


@pytest.mark.parametrize("payload", [{"text": ""}, {"text": "x" * 12001}, {}])
def test_chat_rejects_empty_or_oversized_payload(web, payload):
    assert web.client.post("/api/chat", json=payload).status_code == 422


def test_sessions_are_isolated_and_user_ids_cannot_be_supplied_in_body_or_query(web):
    first_key = session_key(web.client)
    proposal, product = proposal_for(web, web.client)
    web.assistant.store.confirm(first_key, proposal["id"], {product["id"]: product["quantity"]}, explicit=True)
    first_cart = web.client.get("/api/cart").json()
    # The first client owns the shared lifespan; this client only sends requests.
    second = TestClient(main.app, base_url="http://testserver")
    try:
        response = second.get("/api/cart", params={"user_id": first_key, "token": first_cart["token"]})
        assert response.status_code == 200
        assert response.json()["items"] == []
        assert response.json()["token"] != first_cart["token"]
        assert second.cookies.get("ekt_session") != web.client.cookies.get("ekt_session")
        cookie = response.headers["set-cookie"]
        assert "HttpOnly" in cookie and "SameSite=strict" in cookie
        second_key = session_key(second)
        response = second.post("/api/chat", json={"text": "Какая оплата?", "user_id": first_key})
        assert response.status_code == 200
        assert web.assistant.store.history(first_key) == []
        assert web.assistant.store.history(second_key)[0]["content"] == "Какая оплата?"
        assert web.client.get("/api/cart").json()["items"] == first_cart["items"]
        # A read-only public cart token must not double as a write-session cookie.
        second.cookies.clear()
        second.cookies.set("ekt_session", first_cart["token"])
        response = second.get("/api/cart")
        assert response.json()["items"] == []
        assert response.json()["token"] != first_cart["token"]
    finally:
        second.close()


@pytest.mark.parametrize("confirmation", ["да", "ок", "не добавляй", "Да, добавь 3", ""])
def test_confirm_requires_exact_explicit_phrase(web, confirmation):
    proposal, _product = proposal_for(web, web.client)
    response = web.client.post("/api/confirm", json={"proposal_id": proposal["id"], "confirmation": confirmation})
    assert response.status_code == 400
    assert web.client.get("/api/cart").json()["items"] == []
    assert web.refresh.call_count == 0


def test_another_session_cannot_confirm_a_foreign_proposal(web):
    proposal, _product = proposal_for(web, web.client)
    second = TestClient(main.app, base_url="http://testserver")
    try:
        response = second.post("/api/confirm", json={"proposal_id": proposal["id"], "confirmation": "Да, добавь"})
        assert response.status_code == 200
        assert response.json()["status"] == "confirmation_required"
        assert second.get("/api/cart").json()["items"] == []
        assert web.client.get("/api/cart").json()["items"] == []
        assert web.refresh.call_count == 0
    finally:
        second.close()


def test_explicit_confirmation_adds_once_and_link_reads_current_cart(web):
    proposal, product = proposal_for(web, web.client, quantity=2)
    payload = {"proposal_id": proposal["id"], "confirmation": "Да, добавь"}
    response = web.client.post("/api/confirm", json=payload)
    assert response.status_code == 200
    assert response.json()["status"] == "confirmed"
    web.refresh.assert_awaited_once_with(product["id"])
    cart_url = response.json()["cart_url"]
    cart = web.client.get("/api/cart").json()
    assert cart["items"][0]["quantity"] == 2
    assert cart["total"] == product["price"] * 2
    assert cart_url == cart["url"]
    page = web.client.get(cart_url)
    assert page.status_code == 200
    assert html.escape(product["article"]) in page.text
    assert html.escape(product["name"]) in page.text
    assert "2 ×" in page.text
    assert "пока не связаны" in page.text
    assert page.headers["cache-control"] == "no-store"
    assert page.headers["referrer-policy"] == "no-referrer"
    web.client.post("/api/confirm", json=payload)
    assert web.client.get("/api/cart").json()["items"][0]["quantity"] == 2
    web.assistant.store.clear_cart(session_key(web.client))
    current = web.client.get(cart_url)
    assert current.status_code == 200
    assert html.escape(product["article"]) not in current.text
    assert "Итого: 0" in current.text


def test_confirmation_refuses_stock_that_shrank_after_proposal(web):
    proposal, product = proposal_for(web, web.client, quantity=5)
    web.refresh.side_effect = None
    web.refresh.return_value = product | {"quantity": 2, "stock_fresh": True}
    response = web.client.post("/api/confirm", json={"proposal_id": proposal["id"], "confirmation": "Да, добавь"})
    assert response.status_code == 200
    assert response.json()["status"] == "insufficient_stock"
    assert web.client.get("/api/cart").json()["items"] == []


def test_cart_unknown_token_is_404_and_html_is_escaped(web):
    assert web.client.get("/cart/unknown-token").status_code == 404
    name = '<script>alert("catalog")</script>'
    article = '<img src=x onerror="alert(1)">'
    proposal, product = proposal_for(web, web.client, overrides={"name": name, "article": article})
    key = session_key(web.client)
    confirmed = web.assistant.store.confirm(key, proposal["id"], {product["id"]: product["quantity"]}, explicit=True)
    assert confirmed["status"] == "confirmed"
    page = web.client.get(web.client.get("/api/cart").json()["url"])
    assert page.status_code == 200
    assert name not in page.text and article not in page.text
    assert html.escape(name) in page.text and html.escape(article) in page.text


@pytest.mark.parametrize('filename',['spec.xlsx','spec.docx','spec.pdf','photo.jpeg'])
def test_attachment_http_reads_supported_formats_without_mutating_cart(web,monkeypatch,filename):
    from io import BytesIO
    stream=BytesIO()
    if filename.endswith('.xlsx'):
        from openpyxl import Workbook
        document=Workbook()
        document.active.append(['010500006_',2])
        document.save(stream)
    elif filename.endswith('.docx'):
        from docx import Document
        document=Document()
        document.add_paragraph('010500006_ 2')
        document.save(stream)
    elif filename.endswith('.pdf'):
        from reportlab.pdfgen.canvas import Canvas
        document=Canvas(stream)
        document.drawString(20,700,'010500006_ 2')
        document.save()
    else:
        stream.write((ROOT/'дизайн'/'assets'/'product-050200001.jpg').read_bytes())
    reply=AsyncMock(return_value={'text':'Файл прочитан.', 'status':'completed'})
    monkeypatch.setattr(web.assistant,'reply',reply)
    response=web.client.post('/api/attachment',files={'file':(filename,stream.getvalue())},data={'text':'Проверь спецификацию'})
    assert response.status_code==200,response.text
    args=reply.await_args.args
    assert args[1]=='Проверь спецификацию'
    if filename.endswith('.jpeg'):
        assert args[2]['image_data_url'].startswith('data:image/jpeg;base64,')
    else:
        assert '010500006_' in args[2]['text']
    assert web.client.get('/api/cart').json()['items']==[]
    assert response.headers['cache-control']=='no-store'


def test_attachment_http_rejects_invalid_file_and_oversized_caption(web,monkeypatch):
    reply=AsyncMock()
    monkeypatch.setattr(web.assistant,'reply',reply)
    broken=web.client.post('/api/attachment',files={'file':('broken.pdf',b'not pdf')})
    assert broken.status_code==422
    assert 'PDF' in broken.json()['detail']
    oversized=web.client.post('/api/attachment',files={'file':('spec.docx',b'irrelevant')},data={'text':'x'*12001})
    assert oversized.status_code==422
    assert oversized.headers['cache-control']=='no-store'
    reply.assert_not_awaited()


def test_telegram_application_builds_with_installed_api_without_network(web):
    config = replace(web.config, telegram_token="123456789:ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghi")
    assistant = Assistant(web.assistant.catalog, web.assistant.store, config)
    application = build_application(assistant)
    handlers = application.handlers[0]
    commands = {command for handler in handlers if isinstance(handler, CommandHandler) for command in handler.commands}
    assert {"start", "cart", "orders"} <= commands
    assert any(isinstance(handler, CallbackQueryHandler) for handler in handlers)
    assert any(isinstance(handler, MessageHandler) for handler in handlers)
    assert application.error_handlers

    async def close_unstarted_requests():
        # No initialize/start/get_me is called: these would contact Telegram.
        for request in application.bot._request:
            await request.shutdown()

    asyncio.run(close_unstarted_requests())
