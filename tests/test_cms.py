"""Local CMS boundary checks. No production data or external requests."""

from datetime import datetime, timezone
from io import BytesIO
import json
import sqlite3

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.cms import CMSRepository, Filters, create_cms_router
from app.store import Store


@pytest.fixture
def cms(tmp_path, monkeypatch):
    monkeypatch.delenv("CMS_ADMIN_TOKEN", raising=False)
    store = Store(tmp_path / "assistant.sqlite")
    store.add_message("PRIVATE-USER-ID", "user", "PRIVATE-CHAT-CONTENT")
    product = {"id": 1, "article": "DEMO-001", "name": "Кабель", "price": 10, "quantity": 2}
    proposal = store.propose("PRIVATE-USER-ID", [product])
    store.confirm("PRIVATE-USER-ID", proposal["id"], {1: 4}, explicit=True)
    store.save_request("PRIVATE-USER-ID")
    timestamp = datetime(2026, 9, 23, 12, tzinfo=timezone.utc).timestamp()
    for table in ("messages", "proposals", "requests"):
        store._db.execute(f"UPDATE {table} SET created_at=?", (timestamp,))
    rows = [
        {"timestamp": "2026-09-23T00:00:00Z", "seconds": 2, "tools": ["search"], "manager": False,
         "usage": [{"input_tokens": 100, "output_tokens": 10}]},
        {"timestamp": "2026-09-23T23:59:59Z", "seconds": 4, "tools": [], "manager": True,
         "usage": [{"input_tokens": 200, "output_tokens": 20}]},
        {"timestamp": "2026-09-24T00:00:00Z", "seconds": 90, "usage": []},
        {"seconds": 10, "usage": []},
    ]
    (tmp_path / "metrics.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\ninvalid-json\n", encoding="utf-8")
    repository = CMSRepository(tmp_path, tmp_path / "assistant.sqlite")
    app = FastAPI()
    app.state.cms = repository
    app.include_router(create_cms_router())
    with TestClient(app, base_url="http://127.0.0.1", client=("127.0.0.1", 40000)) as client:
        yield client, repository, app
    repository.close()
    store.close()


def metrics(result):
    return {item["key"]: item["value"] for item in result["metrics"]}


def test_access_remote_requires_token_and_does_not_trust_forwarded_header(cms, monkeypatch):
    client, _, app = cms
    assert client.get("/api/admin/summary").status_code == 200
    with TestClient(app, client=("203.0.113.5", 40000)) as remote:
        assert remote.get("/api/admin/summary").status_code == 403
        assert remote.get("/api/admin/summary", headers={"X-Forwarded-For": "127.0.0.1"}).status_code == 403
        monkeypatch.setenv("CMS_ADMIN_TOKEN", "test-token")
        assert remote.get("/api/admin/summary").status_code == 401
        assert remote.get("/api/admin/summary", headers={"X-Admin-Token": "test-token"}).status_code == 200
        assert remote.get("/api/admin/summary", headers={"Authorization": "Bearer test-token"}).status_code == 200
    assert client.get("/api/admin/summary").status_code == 401
    assert client.get("/api/admin/summary", headers={"X-Admin-Token": "wrong"}).status_code == 401
    assert client.get("/api/admin/summary", headers={"X-Admin-Token": "test-token"}).status_code == 200


def test_cross_origin_local_writes_rejected(cms):
    client, _, _ = cms
    response = client.post("/api/admin/tasks", json={"department": "sales", "title": "Тест"},
                           headers={"Origin": "https://outside.example"})
    assert response.status_code == 403
    assert client.get("/api/admin/tasks").json()["tasks"] == []


def test_summary_reads_real_aggregates_and_honours_inclusive_utc_dates(cms):
    client, _, _ = cms
    response = client.get("/api/admin/summary", params={"from": "2026-09-23", "to": "2026-09-23"})
    assert response.status_code == 200
    result = response.json()
    values = metrics(result)
    assert values["sessions"] == values["questions"] == values["proposals"] == values["confirmed_proposals"] == values["saved_requests"] == 1
    assert values["confirmation_rate"] == 100
    assert values["responses"] == 2
    assert values["latency_avg"] == 3
    assert values["latency_p95"] == 4
    assert values["input_tokens"] == 300 and values["output_tokens"] == 30
    assert values["manager_activations"] == 1
    assert values["cost_usd"] is None
    assert result["timeseries"] == [{"date": "2026-09-23", "requests": 2, "errors": 0, "tokens": 330, "latency_ms": 3000}]
    assert any("без даты: 1" in text for text in result["limitations"])
    assert any("пропущено: 1" in text for text in result["limitations"])
    assert "PRIVATE" not in response.text
    assert "DEMO-001" not in response.text
    assert metrics(client.get("/api/admin/summary").json())["responses"] == 4
    empty = metrics(client.get("/api/admin/summary?from=2026-10-01&to=2026-10-01").json())
    assert empty["questions"] == empty["responses"] == 0
    assert empty["latency_avg"] is None


@pytest.mark.parametrize("query", ["department=unknown", "from=oops", "from=2026-09-24&to=2026-09-23", "to=9999-12-31"])
def test_invalid_filters_are_rejected(cms, query):
    assert cms[0].get("/api/admin/summary?" + query).status_code == 422


def test_department_hierarchy_and_metric_filter(cms):
    client, _, _ = cms
    departments = client.get("/api/admin/departments").json()["departments"]
    assert len(departments) == 5
    assert departments[0]["parent_id"] is None
    assert all(row["parent_id"] == "director" for row in departments[1:])
    result = client.get("/api/admin/summary?department=sales").json()
    assert all(item["department"] == "sales" for item in result["metrics"])
    assert metrics(result)["saved_requests"] == 1


def test_task_lifecycle_validation_and_persistence(cms):
    client, repository, _ = cms
    task = {"department": "marketing", "title": "Проверить спрос", "description": "Сравнить обращения", "assignee": "Маркетолог", "due_date": "2026-09-29"}
    response = client.post("/api/admin/tasks", json=task)
    assert response.status_code == 201
    created = response.json()
    assert created["status"] == "todo"
    task_id = created["id"]
    assert client.get("/api/admin/tasks?department=sales").json()["tasks"] == []
    assert len(client.get("/api/admin/tasks?department=marketing").json()["tasks"]) == 1
    assert client.get("/api/admin/tasks?to=2000-01-01").json()["tasks"] == []
    updated = client.patch(f"/api/admin/tasks/{task_id}", json={"status": "done", "due_date": None})
    assert updated.status_code == 200
    assert updated.json()["due_date"] is None
    assert client.get("/api/admin/tasks?status=todo").json()["tasks"] == []
    assert client.get("/api/admin/summary").json()["task_counts"]["done"] == 1
    with sqlite3.connect(repository.data_dir / "cms.sqlite") as persisted:
        assert persisted.execute("SELECT status FROM tasks WHERE id=?", (task_id,)).fetchone()[0] == "done"
    assert client.delete(f"/api/admin/tasks/{task_id}").status_code == 200
    assert client.delete(f"/api/admin/tasks/{task_id}").status_code == 404
    assert client.patch(f"/api/admin/tasks/{task_id}", json={"status": "todo"}).status_code == 404


@pytest.mark.parametrize("body", [{"title": " "}, {"department": "fake"}, {"status": "closed"},
                                   {"status": None}, {"title": None}, {"unexpected": "field"}, {}, {"due_date": "not-a-date"}])
def test_invalid_task_changes_are_rejected(cms, body):
    client, _, _ = cms
    task = client.post("/api/admin/tasks", json={"department": "sales", "title": "Проверка"}).json()
    assert client.patch(f"/api/admin/tasks/{task['id']}", json=body).status_code == 422


def test_content_is_persistent_plain_text(cms):
    client, repository, _ = cms
    updated = {"welcome": "Подберём кабель из каталога", "delivery_note": "Уточните адрес доставки у менеджера."}
    response = client.put("/api/admin/content", json=updated)
    assert response.status_code == 200
    assert all(response.json()[key] == value for key, value in updated.items())
    assert repository.public_content()["welcome"] == updated["welcome"]
    assert client.put("/api/admin/content", json={**updated, "welcome": "<script>bad()</script>"}).status_code == 422
    assert client.put("/api/admin/content", json={**updated, "welcome": " "}).status_code == 422


def test_record_event_discards_chat_and_customer_data(cms):
    client, repository, _ = cms
    repository.record_event("support", "chat_error", {"seconds": 1, "status": "error", "error_type": "TimeoutError",
                                                    "user_id": "PRIVATE", "message": "PRIVATE", "detail": {"secret": "PRIVATE"},
                                                    "source": "contains private text", "cost_usd": float("nan")})
    row = repository._db.execute("SELECT payload FROM events").fetchone()[0]
    assert "PRIVATE" not in row and "source" not in row and "cost_usd" not in row
    assert metrics(client.get("/api/admin/summary?department=support").json())["errors"] == 1
    assert metrics(client.get("/api/admin/summary?department=support&to=2000-01-01").json())["errors"] == 0
    with pytest.raises(ValueError):
        repository.record_event("unknown", "chat_error")


def test_missing_sources_are_not_reported_as_measured_zero(tmp_path):
    repository = CMSRepository(tmp_path, tmp_path / "missing.sqlite")
    try:
        result = repository.summary(Filters())
        values = metrics(result)
        assert result["sources"] == {"assistant_db": "missing", "metrics_log": "missing"}
        assert values["questions"] is None and values["responses"] is None
        assert values["latency_avg"] is None and values["cost_usd"] is None
        assert not (tmp_path / "missing.sqlite").exists()
    finally:
        repository.close()


def test_export_xlsx_typed_values_russian_and_formula_injection(cms):
    client, _, _ = cms
    title = '=HYPERLINK("https://example.invalid")'
    client.post("/api/admin/tasks", json={"department": "sales", "title": title, "assignee": "Менеджер", "due_date": "2026-09-29"})
    response = client.get("/api/admin/exports/xlsx?department=sales")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Показатели", "Задачи"]
    assert book["Показатели"]["A6"].value == "Предложения корзины"
    assert book["Показатели"]["B6"].value == 1
    sheet = book["Задачи"]
    assert sheet["C2"].value == title and sheet["C2"].data_type == "s"
    assert isinstance(sheet["G2"].value, datetime)
    assert sheet["F2"].value == "Менеджер"
    assert sheet.freeze_panes == "A2"
    assert "Tasks" in sheet.tables
    assert "PRIVATE" not in str(list(book["Показатели"].values))


def test_export_pdf_contains_cyrillic_and_filtered_tasks(cms):
    client, _, _ = cms
    client.post("/api/admin/tasks", json={"department": "support", "title": "Проверить сертификаты", "description": "Кабель и электропроводка. " * 50})
    client.post("/api/admin/tasks", json={"department": "sales", "title": "Чужая задача"})
    response = client.get("/api/admin/exports/pdf?department=support")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    pdf = PdfReader(BytesIO(response.content))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Отчёт департаментов" in text
    assert "Проверить сертификаты" in text
    assert "Кабель и электропроводка" in text
    assert "Чужая задача" not in text
    assert "PRIVATE" not in text
    assert all("/Font" in page["/Resources"] for page in pdf.pages)


def test_exports_require_same_access_as_cms(cms, monkeypatch):
    client, _, _ = cms
    monkeypatch.setenv("CMS_ADMIN_TOKEN", "test-token")
    assert client.get("/api/admin/exports/xlsx").status_code == 401
    assert client.get("/api/admin/exports/pdf").status_code == 401
    assert client.get("/api/admin/exports/csv", headers={"X-Admin-Token": "test-token"}).status_code == 422
