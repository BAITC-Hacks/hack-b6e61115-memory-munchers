"""Live local website acceptance run. Only local carts mutate; EKT is read-only."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import time
import os

import httpx
from openpyxl import load_workbook
from pypdf import PdfReader


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--base',default='http://127.0.0.1:8001')
    parser.add_argument('--output',default='data/integration-smoke.json')
    args=parser.parse_args()
    report={'base':args.base,'runs':[]}
    for iteration in range(2):
        with httpx.Client(base_url=args.base,timeout=10) as client:
            run={'iteration':iteration+1,'steps':[]}
            def call(method,path,**kwargs):
                if path.startswith('/api/admin') and os.getenv('CMS_ADMIN_TOKEN'):
                    kwargs['headers']={'X-Admin-Token':os.environ['CMS_ADMIN_TOKEN']}
                started=time.perf_counter()
                result=client.request(method,path,**kwargs)
                elapsed=round(time.perf_counter()-started,3)
                result.raise_for_status()
                run['steps'].append({'path':path.split('?')[0],'status':result.status_code,'seconds':elapsed})
                assert elapsed<8,(path,elapsed)
                return result
            health=call('GET','/health').json()
            assert health['catalog_products']>=15035
            assert call('GET','/').status_code==200
            detail=call('POST','/api/chat',json={'text':'010500006_'}).json()
            product=detail['products'][0]
            assert product['article']=='010500006_'
            assert product.get('certificates') and product.get('properties')
            terms=call('POST','/api/chat',json={'text':'Какие оплата, доставка и минимальная партия?'}).json()
            assert 'Физлица' in terms['text'] and 'упаковки' in terms['text']
            alternatives=call('POST','/api/chat',json={'text':'310100024_'}).json()
            assert len(alternatives['products'])>1
            draft=call('POST','/api/proposal',json={'items':[{'product_id':product['id'],'quantity':2}]}).json()
            assert call('GET','/api/cart').json()['items']==[]
            # An unrelated question must preserve the approved draft's identity.
            call('POST','/api/chat',json={'text':'Какая доставка?'})
            assert call('GET','/api/session').json()['proposal']['id']==draft['id']
            confirmed=call('POST','/api/confirm',json={'proposal_id':draft['id'],'confirmation':'Да, добавь'}).json()
            assert confirmed['status']=='confirmed',confirmed
            cart=call('GET','/api/cart').json()
            assert cart['items'][0]['quantity']==2
            assert call('GET',confirmed['cart_url']).status_code==200
            call('POST','/api/confirm',json={'proposal_id':draft['id'],'confirmation':'Да, добавь'})
            assert call('GET','/api/cart').json()['items'][0]['quantity']==2
            saved=call('POST','/api/orders/save').json()
            assert saved['status']=='request'
            call('DELETE','/api/cart',json={'confirmation':'Да, очистить корзину'})
            repeated=call('POST','/api/orders/'+saved['id']+'/repeat').json()
            assert repeated['status']=='pending'
            assert call('GET','/api/cart').json()['items']==[]
            call('POST','/api/proposal/cancel')
            call('GET','/admin')
            call('GET','/api/admin/summary')
            for fmt in ['xlsx','pdf']:
                data=call('GET','/api/admin/exports/'+fmt).content
                if fmt=='xlsx':
                    assert 'Показатели' in load_workbook(BytesIO(data)).sheetnames
                else:
                    assert 'департамент' in ''.join(p.extract_text() or '' for p in PdfReader(BytesIO(data)).pages).lower()
                output=Path('output/qa')
                output.mkdir(parents=True,exist_ok=True)
                (output/('cms-report.'+fmt)).write_bytes(data)
            run['passed']=True
            run['model']=health['model']
            report['runs'].append(run)
    Path(args.output).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':True,'runs':len(report['runs']),'max_seconds':max(s['seconds'] for r in report['runs'] for s in r['steps'])}))


if __name__=='__main__':
    main()
