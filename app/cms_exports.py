"""In-memory XLSX/PDF exports for the department CMS.

Dependencies: openpyxl>=3.1,<4 and reportlab>=4.2,<5.
PDF requires DejaVu Sans (Linux) or Arial (Windows); CMS_PDF_FONT can override.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import os
from pathlib import Path
import re
from xml.sax.saxutils import escape


DEPARTMENTS = {"all": "Все департаменты", "director": "Директор", "marketing": "Маркетинг",
               "sales": "Продажи", "development": "Разработка", "support": "Поддержка"}
STATUSES = {"todo": "К выполнению", "in_progress": "В работе", "blocked": "Блокировка", "done": "Выполнено"}


def _clean(value):
    # Spreadsheet and XML control characters must never make export fail.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def _period(report):
    filters = report["filters"]
    return f"{DEPARTMENTS.get(filters['department'], filters['department'])}; {filters.get('from') or 'начало истории'} - {filters.get('to') or 'сейчас'}; UTC"


def export_xlsx(report: dict, tasks: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    book = Workbook()
    summary = book.active
    summary.title = "Показатели"
    summary.append(["ЕКТ / Отчёт департаментов"])
    summary.append([_period(report)])
    summary.append(["Сформирован", datetime.fromisoformat(report["generated_at"]).replace(tzinfo=None)])
    summary["B3"].number_format = "dd.mm.yyyy hh:mm:ss"
    summary.append([])
    summary.append(["Показатель", "Значение", "Единица", "Департамент", "Источник", "Примечание"])
    for metric in report["metrics"]:
        summary.append([metric["label"], metric["value"] if metric["value"] is not None else "Нет данных", metric["unit"],
                        DEPARTMENTS[metric["department"]], metric["source"], metric.get("note", "")])
    last_metric = summary.max_row
    summary.append([])
    for note in report["limitations"]:
        summary.append([note])
        summary.merge_cells(start_row=summary.max_row, start_column=1, end_row=summary.max_row, end_column=6)
        summary.row_dimensions[summary.max_row].height = 30
    tasks_sheet = book.create_sheet("Задачи")
    tasks_sheet.append(["ID", "Департамент", "Задача", "Описание", "Статус", "Ответственный", "Дедлайн", "Создана (UTC)", "Изменена (UTC)"])
    for task in tasks:
        tasks_sheet.append([task["id"], DEPARTMENTS[task["department"]], task["title"], task["description"],
                            STATUSES[task["status"]], task["assignee"],
                            date.fromisoformat(task["due_date"]) if task.get("due_date") else None,
                            datetime.fromisoformat(task["created_at"]).replace(tzinfo=None),
                            datetime.fromisoformat(task["updated_at"]).replace(tzinfo=None)])
    for sheet, header, widths in ((summary, 5, [44, 19, 12, 20, 24, 66]),
                                   (tasks_sheet, 1, [10, 20, 44, 70, 20, 24, 16, 23, 23])):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = f"A{header + 1}"
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _clean(cell.value)
                    # Never execute an administrator's task title as an Excel formula.
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="242424")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.###"
        for cell in sheet[header]:
            cell.fill = PatternFill("solid", fgColor="08080A")
            cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        sheet.row_dimensions[header].height = 30
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        last = last_metric if sheet is summary else sheet.max_row
        if last > header:
            table = Table(displayName="Metrics" if sheet is summary else "Tasks", ref=f"A{header}:{get_column_letter(len(widths))}{last}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for row_number in range(header + 1, last + 1):
            sheet.row_dimensions[row_number].height = 48 if sheet is summary else 62
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f"{header}:{header}"
    summary["A1"].font = Font(name="Calibri", size=19, bold=True, color="FF2A1C")
    summary.row_dimensions[1].height = 32
    summary.merge_cells("A1:F1")
    summary.merge_cells("A2:F2")
    summary.row_dimensions[2].height = 28
    for row in tasks_sheet.iter_rows(min_row=2):
        row[6].number_format = "dd.mm.yyyy"
        row[7].number_format = row[8].number_format = "dd.mm.yyyy hh:mm:ss"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name = "EKTCyrillic"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = [os.getenv("CMS_PDF_FONT", ""), "C:/Windows/Fonts/arial.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                  "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
                  "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    for value in candidates:
        path = Path(value) if value else None
        if path and path.is_file():
            pdfmetrics.registerFont(TTFont(name, str(path)))
            return name
    raise RuntimeError("PDF требует шрифт DejaVu Sans/Arial или CMS_PDF_FONT.")


def export_pdf(report: dict, tasks: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font = _pdf_font()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm,
                                 topMargin=17 * mm, bottomMargin=18 * mm,
                                 title="ЕКТ - Отчёт департаментов", author="ЕКТ CMS")
    styles = {
        "title": ParagraphStyle("title", fontName=font, fontSize=21, leading=27, textColor=colors.HexColor("#FF2A1C"), spaceAfter=10),
        "heading": ParagraphStyle("heading", fontName=font, fontSize=14, leading=19, spaceBefore=14, spaceAfter=8),
        "body": ParagraphStyle("body", fontName=font, fontSize=9, leading=13, spaceAfter=5, splitLongWords=True),
        "small": ParagraphStyle("small", fontName=font, fontSize=8, leading=11, spaceAfter=4, textColor=colors.HexColor("#555555")),
        "head": ParagraphStyle("head", fontName=font, fontSize=9, leading=12, textColor=colors.white),
    }

    def paragraph(value, style="body"):
        return Paragraph(escape(_clean(value)).replace("\n", "<br/>"), styles[style])

    story = [paragraph("ЕКТ / Отчёт департаментов", "title"), paragraph(_period(report)),
             paragraph("Сформирован: " + report["generated_at"], "small"), Spacer(1, 6 * mm)]
    rows = [[paragraph(label, "head") for label in ("Показатель", "Результат", "Департамент")]]
    for metric in report["metrics"]:
        value = "Нет данных" if metric["value"] is None else str(metric["value"])
        rows.append([paragraph(metric["label"]), paragraph(value + (" " + metric["unit"] if metric["value"] is not None else "")),
                     paragraph(DEPARTMENTS[metric["department"]])])
    table = LongTable(rows, colWidths=[90 * mm, 39 * mm, 47 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#08080A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F1F1F1"), colors.white]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)
    story.append(paragraph("Источники и границы отчёта", "heading"))
    for metric in report["metrics"]:
        story.append(paragraph(f"{metric['label']}: {metric['source']}. {metric.get('note', '')}", "small"))
    for note in report["limitations"]:
        story.append(paragraph(note, "small"))
    story.append(paragraph(f"Задачи ({len(tasks)})", "heading"))
    if not tasks:
        story.append(paragraph("За выбранный период задач нет."))
    for task in tasks:
        story.append(paragraph(f"#{task['id']} / {task['title']}", "heading"))
        story.append(paragraph(f"{DEPARTMENTS[task['department']]} / {STATUSES[task['status']]}"))
        story.append(paragraph(f"Ответственный: {task['assignee'] or 'не назначен'}. Дедлайн: {task.get('due_date') or 'не задан'}."))
        if task["description"]:
            story.append(paragraph(task["description"]))
        story.append(paragraph(f"Создана: {task['created_at']}. Изменена: {task['updated_at']}.", "small"))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(17 * mm, 10 * mm, "ЕКТ / CMS / Агрегированные данные")
        canvas.drawRightString(193 * mm, 10 * mm, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
